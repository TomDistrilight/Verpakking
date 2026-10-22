# Maakt de voorbeeldbestanden voor de Excel-import (artikelen en leverdata) in public/.
# Gebruik: python3 scripts/maak_voorbeelden.py public  (vereist openpyxl)
import datetime
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PUBLIC = sys.argv[1]

KOP_FONT = Font(bold=True)
KOP_FILL = PatternFill(fill_type='solid', start_color='DCE8F5', end_color='DCE8F5')
TITEL_FONT = Font(bold=True, size=13)
VAST = datetime.datetime(2026, 10, 9, 12, 0, 0)


def nieuw_boek():
    wb = Workbook()
    wb.properties.creator = 'Verpakking'
    wb.properties.lastModifiedBy = 'Verpakking'
    wb.properties.created = VAST
    wb.properties.modified = VAST
    return wb


def kopregel(ws, koppen):
    """Kopregel: vet, lichte vulling, eerste rij vastgezet. Breedtes en opmaak per kolom: (kop, breedte, getalnotatie)."""
    for i, (kop, breedte, notatie) in enumerate(koppen, start=1):
        c = ws.cell(row=1, column=i, value=kop)
        c.font = KOP_FONT
        c.fill = KOP_FILL
        c.alignment = Alignment(vertical='center')
        letter = c.column_letter
        ws.column_dimensions[letter].width = breedte
        # Opmaak op kolomniveau: geen lege cellen, dus het onaangeroerde voorbeeld bevat geen datarijen.
        if notatie:
            ws.column_dimensions[letter].number_format = notatie
    ws.freeze_panes = 'A2'


def uitleg(wb, titel, kolommen, punten):
    ws = wb.create_sheet('Uitleg')
    ws['A1'] = titel
    ws['A1'].font = TITEL_FONT
    for i, kop in enumerate(['Kolom', 'Kop', 'Wat vul je in'], start=1):
        c = ws.cell(row=3, column=i, value=kop)
        c.font = KOP_FONT
        c.fill = KOP_FILL
    rij = 4
    for kolom, kop, tekst in kolommen:
        ws.cell(row=rij, column=1, value=kolom).alignment = Alignment(horizontal='center', vertical='top')
        ws.cell(row=rij, column=2, value=kop).alignment = Alignment(vertical='top')
        ws.cell(row=rij, column=3, value=tekst).alignment = Alignment(wrap_text=True, vertical='top')
        rij += 1
    rij += 1
    ws.cell(row=rij, column=1, value='Goed om te weten').font = KOP_FONT
    rij += 1
    for p in punten:
        ws.cell(row=rij, column=1, value=f'- {p}')
        rij += 1
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 90
    return ws


KOPPEN_ARTIKELEN = [
    ('Artikelcode', 18, '@'),
    ('Lengte binnendoos (mm)', 24, '0'),
    ('Breedte binnendoos (mm)', 25, '0'),
    ('Hoogte binnendoos (mm)', 24, '0'),
    ('Gewicht binnendoos incl. inhoud (kg)', 36, '0.000'),
    ('Artikelen per binnendoos', 26, '0'),
    ('Omschrijving', 40, None),
    # Kantelrechten (ronde 8, #68): ja of nee.
    ('Kantelbaar (ja/nee)', 21, None),
    ('L mag verticaal (ja/nee)', 25, None),
    ('B mag verticaal (ja/nee)', 25, None),
]

# Ingevulde rijen voor het werkblad Voorbeeld: niet kantelbaar, kantelbaar met alleen L verticaal, en kantelbaar met L en B.
VOORBEELD_RIJEN = [
    ('VOORBEELD-1', 300, 200, 150, 2.5, 4, 'Kop en schotel', 'nee', 'nee', 'nee'),
    ('VOORBEELD-2', 400, 250, 120, 1.25, 6, 'Vaas', 'ja', 'ja', 'nee'),
    ('VOORBEELD-3', 350, 300, 200, 3, None, 'Bordenset', 'ja', 'ja', 'ja'),
]


def artikelen():
    wb = nieuw_boek()
    ws = wb.active
    ws.title = 'Artikelen'
    kopregel(ws, KOPPEN_ARTIKELEN)
    # Alleen een waarschuwing: de app leest ook tekst als "2,5", dus niets blokkeren.
    maat = DataValidation(type='decimal', operator='greaterThan', formula1='0', errorStyle='warning', allow_blank=True, showErrorMessage=True)
    maat.error = 'Vul een getal groter dan 0 in.'
    maat.errorTitle = 'Geen geldige waarde'
    maat.add('B2:E10000')
    aantal = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='1', errorStyle='warning', allow_blank=True, showErrorMessage=True)
    aantal.error = 'Vul een geheel getal van minstens 1 in, of laat de cel leeg (= 2).'
    aantal.errorTitle = 'Geen geldig aantal'
    aantal.add('F2:F10000')
    jaNee = DataValidation(type='list', formula1='"ja,nee"', errorStyle='warning', allow_blank=True, showErrorMessage=True)
    jaNee.error = 'Vul ja of nee in, of laat de cel leeg (= nee).'
    jaNee.errorTitle = 'Geen ja of nee'
    jaNee.add('H2:J10000')
    ws.add_data_validation(maat)
    ws.add_data_validation(aantal)
    ws.add_data_validation(jaNee)

    uitleg(
        wb,
        'Uitleg bij het werkblad Artikelen',
        [
            ('A', 'Artikelcode', 'Het artikelnummer. Verplicht; elk artikelnummer één keer per bestand.'),
            ('B', 'Lengte binnendoos (mm)', 'Lengte van de binnendoos in millimeters. Verplicht, groter dan 0.'),
            ('C', 'Breedte binnendoos (mm)', 'Breedte van de binnendoos in millimeters. Verplicht, groter dan 0.'),
            ('D', 'Hoogte binnendoos (mm)', 'Hoogte van de binnendoos rechtop (zoals hij normaal staat) in millimeters. Verplicht, groter dan 0.'),
            ('E', 'Gewicht binnendoos incl. inhoud (kg)', 'Gewicht van één gevulde binnendoos (doos plus inhoud) in kilogram, bijvoorbeeld 2,5. Verplicht, groter dan 0.'),
            ('F', 'Artikelen per binnendoos', 'Hoeveel artikelen er in één binnendoos zitten, als geheel getal. Mag leeg blijven: leeg = 2.'),
            ('G', 'Omschrijving', 'Korte omschrijving van het artikel. Mag leeg blijven.'),
            (
                'H',
                'Kantelbaar (ja/nee)',
                'Ja als de binnendoos op zijn kant mag staan, anders nee. Mag leeg blijven: leeg = nee. Bij ja moet L of B (of allebei) verticaal mogen.',
            ),
            ('I', 'L mag verticaal (ja/nee)', 'Ja als de binnendoos zo mag staan dat de lengte verticaal is. Alleen bij kantelbaar ja. Leeg = nee.'),
            ('J', 'B mag verticaal (ja/nee)', 'Ja als de binnendoos zo mag staan dat de breedte verticaal is. Alleen bij kantelbaar ja. Leeg = nee.'),
        ],
        [
            'Laat de eerste rij met de kolomnamen staan: de app herkent daaraan de kolommen en de eenheden. Vul vanaf rij 2 één artikel per rij in.',
            'Lengtes in millimeters (mm), gewicht in kilogram (kg) inclusief inhoud.',
            'De app kiest bij het importeren standaard het eerste werkblad (Artikelen).',
            'Een nieuw artikelnummer wordt als nieuw artikel aangemaakt. Een bestaand artikel wordt alleen overschreven als je de wijziging in de app bevestigt.',
            'Voor ja en nee mag ook: j en n, yes en no, x (= ja), 1 en 0, of waar en onwaar.',
            'Laat je de kolom Kantelbaar weg, dan zijn nieuwe artikelen niet kantelbaar en blijven de kantelrechten van bestaande artikelen zoals ze zijn. Je kunt ze ook later per artikel aanpassen op het tabblad Artikelen.',
            'Het werkblad Voorbeeld toont drie ingevulde rijen. Vul je eigen artikelen in op het werkblad Artikelen.',
            'Importeren gaat via het tabblad Excel-import.',
        ],
    )
    vb = wb.create_sheet('Voorbeeld')
    kopregel(vb, KOPPEN_ARTIKELEN)
    for rij in VOORBEELD_RIJEN:
        vb.append(rij)
    wb.save(f'{PUBLIC}/voorbeeld-artikelen.xlsx')


def leverdata():
    wb = nieuw_boek()
    ws = wb.active
    ws.title = 'Leverdata'
    kopregel(ws, [('Artikelcode', 18, '@'), ('Verwachte leverdatum', 24, 'dd-mm-yyyy')])

    uitleg(
        wb,
        'Uitleg bij het werkblad Leverdata',
        [
            ('A', 'Artikelcode', 'Het artikelnummer zoals het in het Overzicht staat.'),
            ('B', 'Verwachte leverdatum', 'De verwachte leverdatum: een datumcel of tekst als 31-12-2026.'),
        ],
        [
            'Laat de eerste rij met de kolomnamen staan en vul vanaf rij 2 één artikel per rij in.',
            'Alleen oplossingen die nog niet fysiek gecontroleerd zijn krijgen de datum; bij een gecontroleerde oplossing blijft de leverdatum staan.',
            'Artikelnummers die niet in het Overzicht staan worden overgeslagen. Er worden geen artikelen of regels aangemaakt.',
            'Weeknummers en levertijden in dagen worden niet herkend; vul een datum in.',
            'Staat een artikelnummer dubbel in het bestand, dan telt alleen de eerste rij.',
            'De app kiest bij het importeren standaard het eerste werkblad (Leverdata).',
            'Importeren gaat via het tabblad Overzicht, knop Leverdata importeren.',
        ],
    )
    wb.save(f'{PUBLIC}/voorbeeld-leverdata.xlsx')


artikelen()
leverdata()
